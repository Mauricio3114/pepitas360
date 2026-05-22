from functools import wraps
from datetime import date, datetime
from io import BytesIO

from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file
from flask_login import login_required, current_user

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

from app.models.movimentacao_financeira import MovimentacaoFinanceira
from app.models.compromisso import Compromisso

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


relatorios_bp = Blueprint(
    "relatorios",
    __name__,
    url_prefix="/relatorios"
)


def requer_relatorios(func):

    @wraps(func)
    def wrapper(*args, **kwargs):

        if not current_user.pode_acessar("relatorios"):
            flash("Você não possui permissão para acessar relatórios.", "danger")
            return redirect(url_for("dashboard.index"))

        return func(*args, **kwargs)

    return wrapper


def datas_filtro():

    hoje = date.today()
    inicio_mes = hoje.replace(day=1)

    data_inicio_str = request.args.get("data_inicio")
    data_fim_str = request.args.get("data_fim")

    if data_inicio_str:
        data_inicio = datetime.strptime(data_inicio_str, "%Y-%m-%d").date()
    else:
        data_inicio = inicio_mes

    if data_fim_str:
        data_fim = datetime.strptime(data_fim_str, "%Y-%m-%d").date()
    else:
        data_fim = hoje

    return data_inicio, data_fim


def buscar_dados(data_inicio, data_fim):

    movimentacoes = MovimentacaoFinanceira.query.filter(
        MovimentacaoFinanceira.empresa_id == current_user.empresa_id,
        MovimentacaoFinanceira.data_movimentacao >= data_inicio,
        MovimentacaoFinanceira.data_movimentacao <= data_fim
    ).order_by(
        MovimentacaoFinanceira.data_movimentacao.asc()
    ).all()

    compromissos = Compromisso.query.filter(
        Compromisso.empresa_id == current_user.empresa_id,
        Compromisso.data_vencimento >= data_inicio,
        Compromisso.data_vencimento <= data_fim
    ).order_by(
        Compromisso.data_vencimento.asc()
    ).all()

    total_entradas = sum(
        item.valor for item in movimentacoes
        if item.tipo == "entrada"
    )

    total_saidas = sum(
        item.valor for item in movimentacoes
        if item.tipo == "saida"
    )

    saldo = total_entradas - total_saidas

    total_pagamentos = sum(
        item.valor for item in compromissos
        if item.tipo in ["pagamento", "boleto", "transferencia", "dinheiro"]
    )

    return movimentacoes, compromissos, total_entradas, total_saidas, saldo, total_pagamentos


@relatorios_bp.route("/")
@login_required
@requer_relatorios
def index():

    data_inicio, data_fim = datas_filtro()

    movimentacoes, compromissos, total_entradas, total_saidas, saldo, total_pagamentos = buscar_dados(
        data_inicio,
        data_fim
    )

    return render_template(
        "relatorios/index.html",
        data_inicio=data_inicio,
        data_fim=data_fim,
        movimentacoes=movimentacoes,
        compromissos=compromissos,
        total_entradas=total_entradas,
        total_saidas=total_saidas,
        saldo=saldo,
        total_pagamentos=total_pagamentos
    )


@relatorios_bp.route("/financeiro/pdf")
@login_required
@requer_relatorios
def financeiro_pdf():

    data_inicio, data_fim = datas_filtro()

    movimentacoes, compromissos, total_entradas, total_saidas, saldo, total_pagamentos = buscar_dados(
        data_inicio,
        data_fim
    )

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    styles = getSampleStyleSheet()
    elementos = []

    cabecalho = Table([
        [
            Paragraph(
                "<font size=22><b>Pepitas360</b></font><br/>"
                "<font size=10 color='#94a3b8'>Sistema Financeiro e Administrativo</font>",
                styles["Normal"]
            ),

            Paragraph(
                f"<font size=10><b>Emitido em:</b><br/>"
                f"{datetime.now().strftime('%d/%m/%Y %H:%M')}</font>",
                styles["Normal"]
            )
        ]
    ], colWidths=[340, 150])

    cabecalho.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#020617")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ("BOX", (0, 0), (-1, -1), 0, colors.white),
        ("LEFTPADDING", (0, 0), (-1, -1), 20),
        ("RIGHTPADDING", (0, 0), (-1, -1), 20),
        ("TOPPADDING", (0, 0), (-1, -1), 18),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 18),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    elementos.append(cabecalho)
    elementos.append(Spacer(1, 22))

    titulo = Paragraph(
        "<font size=18><b>Relatório Financeiro Completo</b></font>",
        styles["Title"]
    )

    periodo = Paragraph(
        f"<font size=11><b>Período:</b> "
        f"{data_inicio.strftime('%d/%m/%Y')} até "
        f"{data_fim.strftime('%d/%m/%Y')}</font>",
        styles["Normal"]
    )

    elementos.append(titulo)
    elementos.append(Spacer(1, 12))
    elementos.append(periodo)
    elementos.append(Spacer(1, 24))

    cards = Table([
        [
            Paragraph(
                f"<font color='white'><b>Entradas</b><br/><br/>"
                f"<font size=16>R$ {total_entradas:.2f}</font></font>",
                styles["BodyText"]
            ),

            Paragraph(
                f"<font color='white'><b>Saídas</b><br/><br/>"
                f"<font size=16>R$ {total_saidas:.2f}</font></font>",
                styles["BodyText"]
            ),

            Paragraph(
                f"<font color='white'><b>Saldo</b><br/><br/>"
                f"<font size=16>R$ {saldo:.2f}</font></font>",
                styles["BodyText"]
            ),
        ]
    ], colWidths=[170, 170, 170])

    cards.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor("#166534")),
        ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#991b1b")),
        ("BACKGROUND", (2, 0), (2, 0), colors.HexColor("#1d4ed8")),
        ("BOX", (0, 0), (-1, -1), 0, colors.white),
        ("LEFTPADDING", (0, 0), (-1, -1), 18),
        ("RIGHTPADDING", (0, 0), (-1, -1), 18),
        ("TOPPADDING", (0, 0), (-1, -1), 18),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 18),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    elementos.append(cards)
    elementos.append(Spacer(1, 24))

    elementos.append(
        Paragraph(
            "<font size=14><b>Movimentações Financeiras</b></font>",
            styles["Heading2"]
        )
    )

    elementos.append(Spacer(1, 10))

    dados_mov = [["Data", "Tipo", "Descrição", "Categoria", "Valor"]]

    for item in movimentacoes:
        dados_mov.append([
            item.data_movimentacao.strftime("%d/%m/%Y"),
            item.tipo,
            item.descricao,
            item.categoria or "-",
            f"R$ {item.valor:.2f}"
        ])

    if len(dados_mov) == 1:
        dados_mov.append(["-", "-", "Nenhuma movimentação encontrada", "-", "-"])

    tabela_mov = Table(dados_mov, colWidths=[70, 70, 170, 90, 80])

    tabela_mov.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 10),
        ("TOPPADDING", (0, 0), (-1, 0), 10),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.black),
        ("PADDING", (0, 1), (-1, -1), 7),
    ]))

    elementos.append(tabela_mov)
    elementos.append(Spacer(1, 24))

    elementos.append(
        Paragraph(
            "<font size=14><b>Pagamentos, Agenda e Compromissos</b></font>",
            styles["Heading2"]
        )
    )

    elementos.append(Spacer(1, 10))

    dados_comp = [["Data", "Tipo", "Título", "Status", "Valor"]]

    for item in compromissos:
        dados_comp.append([
            item.data_vencimento.strftime("%d/%m/%Y"),
            item.tipo,
            item.titulo,
            item.status,
            f"R$ {item.valor:.2f}"
        ])

    if len(dados_comp) == 1:
        dados_comp.append(["-", "-", "Nenhum compromisso encontrado", "-", "-"])

    tabela_comp = Table(dados_comp, colWidths=[70, 70, 170, 80, 80])

    tabela_comp.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 10),
        ("TOPPADDING", (0, 0), (-1, 0), 10),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.black),
        ("PADDING", (0, 1), (-1, -1), 7),
    ]))

    elementos.append(tabela_comp)
    elementos.append(Spacer(1, 25))

    rodape = Paragraph(
        "<font size=9 color='#64748b'>"
        "Relatório gerado automaticamente pelo Pepitas360"
        "</font>",
        styles["Normal"]
    )

    elementos.append(rodape)

    doc.build(elementos)

    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="relatorio_financeiro_pepitas360.pdf",
        mimetype="application/pdf"
    )


@relatorios_bp.route("/financeiro/excel")
@login_required
@requer_relatorios
def financeiro_excel():

    data_inicio, data_fim = datas_filtro()

    movimentacoes, compromissos, total_entradas, total_saidas, saldo, total_pagamentos = buscar_dados(
        data_inicio,
        data_fim
    )

    workbook = Workbook()

    aba_resumo = workbook.active
    aba_resumo.title = "Resumo"

    header_fill = PatternFill(
        start_color="111827",
        end_color="111827",
        fill_type="solid"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    titulo_font = Font(
        bold=True,
        size=14
    )

    aba_resumo["A1"] = "Pepitas360 - Relatório Financeiro"
    aba_resumo["A1"].font = titulo_font

    aba_resumo["A3"] = "Período"
    aba_resumo["B3"] = f"{data_inicio.strftime('%d/%m/%Y')} até {data_fim.strftime('%d/%m/%Y')}"

    resumo = [
        ("Entradas", total_entradas),
        ("Saídas", total_saidas),
        ("Saldo", saldo),
        ("Pagamentos/Compromissos", total_pagamentos),
    ]

    linha = 5

    for titulo, valor in resumo:

        aba_resumo[f"A{linha}"] = titulo
        aba_resumo[f"B{linha}"] = valor

        linha += 1

    aba_resumo.column_dimensions["A"].width = 32
    aba_resumo.column_dimensions["B"].width = 22

    aba_mov = workbook.create_sheet("Movimentações")

    headers_mov = [
        "Data",
        "Tipo",
        "Descrição",
        "Categoria",
        "Valor"
    ]

    for col, header in enumerate(headers_mov, 1):

        cell = aba_mov.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font

    linha = 2

    for item in movimentacoes:

        aba_mov.cell(
            row=linha,
            column=1,
            value=item.data_movimentacao.strftime("%d/%m/%Y")
        )

        aba_mov.cell(
            row=linha,
            column=2,
            value=item.tipo
        )

        aba_mov.cell(
            row=linha,
            column=3,
            value=item.descricao
        )

        aba_mov.cell(
            row=linha,
            column=4,
            value=item.categoria or "-"
        )

        aba_mov.cell(
            row=linha,
            column=5,
            value=float(item.valor)
        )

        linha += 1

    tamanhos_mov = [16, 16, 40, 24, 18]

    for i, tamanho in enumerate(tamanhos_mov, 1):
        aba_mov.column_dimensions[get_column_letter(i)].width = tamanho

    aba_comp = workbook.create_sheet("Compromissos")

    headers_comp = [
        "Data",
        "Tipo",
        "Título",
        "Status",
        "Valor"
    ]

    for col, header in enumerate(headers_comp, 1):

        cell = aba_comp.cell(row=1, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font

    linha = 2

    for item in compromissos:

        aba_comp.cell(
            row=linha,
            column=1,
            value=item.data_vencimento.strftime("%d/%m/%Y")
        )

        aba_comp.cell(
            row=linha,
            column=2,
            value=item.tipo
        )

        aba_comp.cell(
            row=linha,
            column=3,
            value=item.titulo
        )

        aba_comp.cell(
            row=linha,
            column=4,
            value=item.status
        )

        aba_comp.cell(
            row=linha,
            column=5,
            value=float(item.valor)
        )

        linha += 1

    tamanhos_comp = [16, 16, 40, 20, 18]

    for i, tamanho in enumerate(tamanhos_comp, 1):
        aba_comp.column_dimensions[get_column_letter(i)].width = tamanho

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="relatorio_financeiro_pepitas360.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )